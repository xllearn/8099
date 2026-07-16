from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document

from app import main as main_module
from app.formal_body_safety import publish_docx_atomically
from app.generation.base import ReportGenerationResult


EXPECTED_TIMING_FIELDS = {
    "prepare_ms",
    "attachment_download_ms",
    "attachment_parse_ms",
    "evidence_build_ms",
    "compact_ms",
    "dify_ms",
    "generation_ms",
    "repair_ms",
    "quality_gate_ms",
    "local_quality_gate_ms",
    "word_render_ms",
    "word_scan_ms",
    "word_publish_ms",
    "export_check_ms",
    "total_ms",
}


def _running_record(run_id: str, pack_id: str) -> dict:
    return {
        "success": True,
        "run_id": run_id,
        "pack_id": pack_id,
        "status": "running",
        "workflow_run_id": "",
        "report_title": "",
        "report_markdown": "",
        "quality_check": {"passed": None, "issues": []},
        "generation_warnings": [],
        "warnings": [],
        "remaining_issues": [],
        "version": 1,
    }


def _pack(pack_id: str) -> dict:
    return {
        "pack_id": pack_id,
        "primary_materials": [
            {
                "material_role": "primary",
                "menu_code": "project_notice",
                "articleid": "article-1",
                "title": "项目公告",
                "content_text": "公告明确申报时间和采购范围。" * 80,
                "attachments": [],
            }
        ],
        "auxiliary_materials": [],
        "warnings": [],
    }


class S3PipelineTimingTests(unittest.TestCase):
    def test_run_timing_schema_covers_every_s3_stage(self) -> None:
        self.assertTrue(EXPECTED_TIMING_FIELDS.issubset(main_module.RUN_DEFAULT_TIMINGS))

    def test_prepare_records_complete_non_negative_timings(self) -> None:
        row = {
            "menu_code": "project_notice",
            "articleid": "article-1",
            "title": "项目公告",
            "content": "公告明确申报时间和采购范围。",
        }
        request = main_module.SelectionPreviewRequest(
            primary_materials=[main_module.MaterialRef(menu_code="project_notice", articleid="article-1")],
            enable_attachment_download=False,
        )
        with patch.object(main_module, "_fetch_database_material_rows", return_value={("project_notice", "article-1"): row}), patch.object(
            main_module, "_fetch_database_attachments", return_value={}
        ), patch.object(main_module, "_write_database_evidence_pack"), patch.dict(
            os.environ,
            {"ENABLE_VBP_FACT_EXTRACTION": "false"},
            clear=False,
        ):
            result = main_module._build_database_evidence_pack(request)

        timings = result["timings"]
        self.assertTrue(EXPECTED_TIMING_FIELDS.issubset(timings))
        self.assertTrue(all(isinstance(value, int) and value >= 0 for value in timings.values()))

    def test_success_and_provider_failure_runs_persist_complete_timings(self) -> None:
        scenarios = (
            ("success", None),
            ("provider_failure", main_module.DifyWorkflowError("TIMEOUT", "timeout", "timeout")),
        )
        for name, failure in scenarios:
            with self.subTest(name=name), tempfile.TemporaryDirectory() as tmpdir:
                root = Path(tmpdir)
                run_id = f"run_{name}_case"
                pack_id = f"pack_{name}_case"
                generator = Mock()
                if failure is not None:
                    generator.generate.side_effect = failure
                else:
                    generator.generate.return_value = ReportGenerationResult(
                        success=True,
                        provider="dify",
                        provider_run_id="workflow-1",
                        report_title="项目公告分析",
                        report_markdown="## 核心结论\n\n公告明确申报时间和采购范围。",
                        quality_check={"passed": True, "issues": []},
                        metadata={"status": "finished", "workflow_run_id": "workflow-1"},
                    )
                with patch.object(main_module, "_analysis_run_dir", return_value=root / "runs", create=True), patch.object(
                    main_module, "_database_evidence_pack_dir", return_value=root / "packs", create=True
                ), patch.object(main_module, "_configured_report_generator", return_value=generator):
                    main_module._write_database_evidence_pack(_pack(pack_id))
                    main_module._write_analysis_run(_running_record(run_id, pack_id))
                    main_module._execute_analysis_run_background(pack_id, run_id)
                    saved = main_module._read_analysis_run(run_id)

                timings = saved["timings"]
                self.assertTrue(EXPECTED_TIMING_FIELDS.issubset(timings))
                self.assertTrue(all(isinstance(value, int) and value >= 0 for value in timings.values()))

    def test_timing_persistence_failure_does_not_block_terminal_run(self) -> None:
        with patch.object(main_module, "_read_analysis_run", return_value={"run_id": "run-1", "status": "finished"}), patch.object(
            main_module, "_write_analysis_run", side_effect=OSError("disk unavailable")
        ):
            main_module._persist_run_timings_nonblocking(
                "run-1",
                {"dify_ms": 5, "total_ms": 5},
            )

    def test_atomic_word_publish_records_render_scan_and_publish(self) -> None:
        events: list[tuple[str, int]] = []
        with tempfile.TemporaryDirectory() as tmpdir:
            destination = Path(tmpdir) / "report.docx"

            def render(path: Path) -> None:
                document = Document()
                document.add_paragraph("公告明确申报时间和采购范围。")
                document.save(path)

            publish_docx_atomically(destination, render, timing_observer=lambda stage, elapsed: events.append((stage, elapsed)))

        self.assertEqual([stage for stage, _ in events], ["word_render_ms", "word_scan_ms", "word_publish_ms"])
        self.assertTrue(all(isinstance(elapsed, int) and elapsed >= 0 for _, elapsed in events))

    def test_timing_observer_failure_does_not_block_word_publish(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            destination = Path(tmpdir) / "report.docx"

            def render(path: Path) -> None:
                document = Document()
                document.add_paragraph("公告明确申报时间和采购范围。")
                document.save(path)

            def fail_observer(stage: str, elapsed: int) -> None:
                raise OSError(f"timing sink unavailable: {stage}:{elapsed}")

            publish_docx_atomically(destination, render, timing_observer=fail_observer)
            self.assertTrue(destination.is_file())


if __name__ == "__main__":
    unittest.main()
