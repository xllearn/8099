from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from app.formal_body import FormalBodyDocument
from app.formal_body_safety import (
    FORBIDDEN_PHRASES,
    FormalBodySafetyError,
    publish_docx_atomically,
    sanitize_formal_body,
    scan_docx,
    scan_formal_body,
)


def _report_ir_with_text(text: str) -> dict:
    return {
        "title": "项目公告分析",
        "lead_paragraphs": [text],
        "sections": [],
        "enterprise_tips": [],
        "disclaimer": "",
    }


class FormalBodySafetyTests(unittest.TestCase):
    def test_all_13_phrases_are_detected_in_four_formal_carriers(self) -> None:
        self.assertEqual(len(FORBIDDEN_PHRASES), 13)
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            for phrase in FORBIDDEN_PHRASES:
                with self.subTest(phrase=phrase, carrier="markdown"):
                    result = scan_formal_body(FormalBodyDocument(markdown=f"# 标题\n\n正式内容：{phrase}。"))
                    self.assertFalse(result.safe)
                    self.assertIn(phrase, {hit.phrase for hit in result.hits})

                with self.subTest(phrase=phrase, carrier="report_ir"):
                    result = scan_formal_body(FormalBodyDocument(report_ir=_report_ir_with_text(f"正式内容：{phrase}。")))
                    self.assertFalse(result.safe)
                    self.assertIn(phrase, {hit.phrase for hit in result.hits})

                with self.subTest(phrase=phrase, carrier="docx_paragraph"):
                    path = root / f"paragraph-{FORBIDDEN_PHRASES.index(phrase)}.docx"
                    doc = Document()
                    doc.add_paragraph(f"正式内容：{phrase}。")
                    doc.save(path)
                    hits = scan_docx(path)
                    self.assertIn(phrase, {hit.phrase for hit in hits})

                with self.subTest(phrase=phrase, carrier="docx_table"):
                    path = root / f"table-{FORBIDDEN_PHRASES.index(phrase)}.docx"
                    doc = Document()
                    table = doc.add_table(rows=1, cols=1)
                    table.cell(0, 0).text = f"正式内容：{phrase}。"
                    doc.save(path)
                    hits = scan_docx(path)
                    self.assertIn(phrase, {hit.phrase for hit in hits})

    def test_nfkc_and_whitespace_variants_are_detected(self) -> None:
        variants = ["需　人工 核验", "待\t确\n认", "请\u00a0核验"]
        for value in variants:
            with self.subTest(value=value):
                result = scan_formal_body(FormalBodyDocument(markdown=f"# 标题\n\n{value}"))
                self.assertFalse(result.safe)

    def test_docx_scan_covers_cross_run_nested_table_header_and_footer(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "all-carriers.docx"
            doc = Document()
            paragraph = doc.add_paragraph()
            paragraph.add_run("需人工")
            paragraph.add_run("核验")
            outer = doc.add_table(rows=1, cols=1)
            nested = outer.cell(0, 0).add_table(rows=1, cols=1)
            nested.cell(0, 0).text = "资料未显示"
            doc.sections[0].header.paragraphs[0].text = "请核验"
            doc.sections[0].footer.paragraphs[0].text = "待确认"
            doc.save(path)

            hits = scan_docx(path)

        self.assertEqual(
            {hit.phrase for hit in hits},
            {"需人工核验", "资料未显示", "请核验", "待确认"},
        )
        self.assertTrue(any("nested_table" in hit.location for hit in hits))
        self.assertTrue(any("header" in hit.location for hit in hits))
        self.assertTrue(any("footer" in hit.location for hit in hits))

    def test_clean_formal_body_has_zero_false_positives(self) -> None:
        document = FormalBodyDocument(
            markdown="# 项目公告分析\n\n公告明确了申报时间、产品范围和执行要求。",
            report_ir=_report_ir_with_text("企业应按照公告列明的时间提交材料。"),
        )

        result = scan_formal_body(document)

        self.assertTrue(result.safe)
        self.assertTrue(result.has_body)
        self.assertEqual(result.hits, ())

    def test_title_without_body_is_not_formal_body(self) -> None:
        for document in (
            FormalBodyDocument(markdown="# 只有标题"),
            FormalBodyDocument(report_ir={"title": "只有标题", "sections": []}),
        ):
            with self.subTest(document=document):
                result = scan_formal_body(document)
                self.assertFalse(result.has_body)
                self.assertFalse(result.safe)

    def test_sanitizer_only_deletes_unsafe_segments_and_does_not_add_facts(self) -> None:
        original = FormalBodyDocument(
            markdown="# 项目公告分析\n\n公告明确申报截止时间为6月30日。\n\n该结论需人工核验。",
            report_ir={
                "title": "项目公告分析",
                "lead_paragraphs": ["公告明确申报截止时间为6月30日。", "资料未显示其他时间。"],
                "sections": [
                    {
                        "heading": "执行要求",
                        "paragraphs": ["企业应按公告提交材料。"],
                        "highlights": ["请核验"],
                        "tables": [],
                    }
                ],
                "enterprise_tips": [],
                "disclaimer": "",
            },
        )

        result = sanitize_formal_body(original)

        self.assertTrue(result.safe)
        self.assertTrue(result.has_body)
        self.assertEqual(result.hits, ())
        self.assertIn("公告明确申报截止时间为6月30日。", result.document.markdown)
        self.assertNotIn("需人工核验", result.document.markdown)
        self.assertEqual(result.document.report_ir["lead_paragraphs"], ["公告明确申报截止时间为6月30日。"])
        self.assertEqual(result.document.report_ir["sections"][0]["paragraphs"], ["企业应按公告提交材料。"])
        self.assertEqual(result.document.report_ir["sections"][0]["highlights"], [])
        self.assertNotIn("evidence_pack", result.document.report_ir)

    def test_atomic_publish_scans_before_replacing_destination(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            destination = Path(tmpdir) / "report.docx"

            def render(path: Path) -> None:
                doc = Document()
                doc.add_paragraph("公告明确了申报时间。")
                doc.save(path)

            publish_docx_atomically(destination, render)

            self.assertTrue(destination.exists())
            self.assertEqual(scan_docx(destination), ())
            self.assertEqual(list(Path(tmpdir).glob(".word-staging-*")), [])

    def test_unsafe_docx_is_deleted_without_publishing_destination(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            destination = Path(tmpdir) / "report.docx"

            def render(path: Path) -> None:
                doc = Document()
                doc.add_paragraph("该内容需人工复核。")
                doc.save(path)

            with self.assertRaises(FormalBodySafetyError):
                publish_docx_atomically(destination, render)

            self.assertFalse(destination.exists())
            self.assertEqual(list(Path(tmpdir).rglob("*.docx")), [])

    def test_atomic_replace_failure_removes_staging_file_and_destination(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            destination = Path(tmpdir) / "report.docx"

            def render(path: Path) -> None:
                doc = Document()
                doc.add_paragraph("公告明确了执行要求。")
                doc.save(path)

            with patch("app.formal_body_safety.os.replace", side_effect=OSError("replace failed")):
                with self.assertRaises(FormalBodySafetyError):
                    publish_docx_atomically(destination, render)

            self.assertFalse(destination.exists())
            self.assertEqual(list(Path(tmpdir).rglob("*.docx")), [])

    def test_cleanup_failure_revokes_published_destination_and_retries_cleanup(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            destination = Path(tmpdir) / "report.docx"

            def render(path: Path) -> None:
                doc = Document()
                doc.add_paragraph("公告明确了执行要求。")
                doc.save(path)

            real_rmtree = __import__("shutil").rmtree
            calls = 0

            def fail_once(path: Path) -> None:
                nonlocal calls
                calls += 1
                if calls == 1:
                    raise OSError("cleanup failed")
                real_rmtree(path)

            with patch("app.formal_body_safety.shutil.rmtree", side_effect=fail_once):
                with self.assertRaises(FormalBodySafetyError):
                    publish_docx_atomically(destination, render)

            self.assertFalse(destination.exists())
            self.assertEqual(list(Path(tmpdir).rglob("*.docx")), [])
            self.assertEqual(list(Path(tmpdir).glob(".word-staging-*")), [])


if __name__ == "__main__":
    unittest.main()
