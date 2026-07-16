from __future__ import annotations

import copy
import logging
import os
import re
import shutil
import tempfile
import time
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Iterable

from docx import Document

from app.formal_body import FormalBodyDocument, REPORT_IR_METADATA_FIELDS


logger = logging.getLogger("medical_notice_analyzer.formal_body_safety")


FORBIDDEN_PHRASES = (
    "原文未披露",
    "需人工核验",
    "需人工复核",
    "证据不足",
    "无法确认",
    "请核验",
    "建议人工确认",
    "资料未显示",
    "未在原文中找到",
    "根据有限信息",
    "以上内容需复核",
    "待确认",
    "待核实",
)

_SENTENCE_SPLIT_RE = re.compile(r"(?<=[。！？!?；;])")


class FormalBodySafetyError(RuntimeError):
    pass


@dataclass(frozen=True)
class ForbiddenPhraseHit:
    phrase: str
    location: str


@dataclass(frozen=True)
class FormalBodySafetyResult:
    document: FormalBodyDocument
    safe: bool
    has_body: bool
    hits: tuple[ForbiddenPhraseHit, ...]
    removed_phrases: tuple[str, ...] = ()


def normalize_for_match(value: Any) -> str:
    normalized = unicodedata.normalize("NFKC", str(value or ""))
    return re.sub(r"\s+", "", normalized)


def find_forbidden_phrases(value: Any, location: str) -> tuple[ForbiddenPhraseHit, ...]:
    normalized = normalize_for_match(value)
    return tuple(
        ForbiddenPhraseHit(phrase=phrase, location=location)
        for phrase in FORBIDDEN_PHRASES
        if normalize_for_match(phrase) in normalized
    )


def scan_formal_body(document: FormalBodyDocument) -> FormalBodySafetyResult:
    hits: list[ForbiddenPhraseHit] = []
    for location, text in document.text_segments():
        hits.extend(find_forbidden_phrases(text, location))
    deduped = _dedupe_hits(hits)
    has_body = document.has_content()
    return FormalBodySafetyResult(
        document=document,
        safe=bool(has_body and not deduped),
        has_body=has_body,
        hits=deduped,
    )


def sanitize_formal_body(document: FormalBodyDocument) -> FormalBodySafetyResult:
    original = scan_formal_body(document)
    if not original.hits:
        return original
    removed: list[str] = []
    markdown = _sanitize_markdown(document.markdown, removed)
    report_ir = _sanitize_report_ir(document.report_ir, removed)
    sanitized = FormalBodyDocument(markdown=markdown, report_ir=report_ir)
    scanned = scan_formal_body(sanitized)
    return FormalBodySafetyResult(
        document=sanitized,
        safe=scanned.safe,
        has_body=scanned.has_body,
        hits=scanned.hits,
        removed_phrases=tuple(dict.fromkeys(removed)),
    )


def scan_docx(path: Path) -> tuple[ForbiddenPhraseHit, ...]:
    try:
        document = Document(path)
    except Exception as exc:  # noqa: BLE001
        raise FormalBodySafetyError("DOCX 扫描失败") from exc

    hits: list[ForbiddenPhraseHit] = []
    hits.extend(_scan_paragraphs(document.paragraphs, "docx.body"))
    for table_index, table in enumerate(document.tables):
        hits.extend(_scan_table(table, f"docx.table[{table_index}]", depth=0))

    for section_index, section in enumerate(document.sections):
        parts = (
            ("header", section.header),
            ("first_page_header", section.first_page_header),
            ("even_page_header", section.even_page_header),
            ("footer", section.footer),
            ("first_page_footer", section.first_page_footer),
            ("even_page_footer", section.even_page_footer),
        )
        for part_name, part in parts:
            prefix = f"docx.section[{section_index}].{part_name}"
            hits.extend(_scan_paragraphs(part.paragraphs, prefix))
            for table_index, table in enumerate(part.tables):
                hits.extend(_scan_table(table, f"{prefix}.table[{table_index}]", depth=0))
    return _dedupe_hits(hits)


def _observe_timing(observer: Callable[[str, int], None] | None, stage: str, started_ns: int) -> None:
    if observer is None:
        return
    elapsed_ms = max(0, time.monotonic_ns() - started_ns) // 1_000_000
    try:
        observer(stage, elapsed_ms)
    except Exception as exc:  # noqa: BLE001
        logger.warning("word_timing_observer_failed stage=%s error_type=%s", stage, exc.__class__.__name__)


def publish_docx_atomically(
    destination: Path,
    render: Callable[[Path], None],
    *,
    timing_observer: Callable[[str, int], None] | None = None,
) -> None:
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    staging_dir = Path(tempfile.mkdtemp(prefix=".word-staging-", dir=destination.parent))
    staging_path = staging_dir / "candidate.docx"
    published = False
    try:
        render(staging_path)
        if not staging_path.is_file():
            raise FormalBodySafetyError("DOCX 临时文件未生成")
        scan_started_ns = time.monotonic_ns()
        try:
            hits = scan_docx(staging_path)
        finally:
            _observe_timing(timing_observer, "word_scan_ms", scan_started_ns)
        if hits:
            phrases = ",".join(dict.fromkeys(hit.phrase for hit in hits))
            raise FormalBodySafetyError(f"DOCX 正文安全扫描未通过:{phrases}")
        publish_started_ns = time.monotonic_ns()
        try:
            os.replace(staging_path, destination)
        finally:
            _observe_timing(timing_observer, "word_publish_ms", publish_started_ns)
        published = True
    except FormalBodySafetyError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise FormalBodySafetyError("DOCX 原子发布失败") from exc
    finally:
        cleanup_error: Exception | None = None
        try:
            if staging_path.exists():
                staging_path.unlink()
            shutil.rmtree(staging_dir)
        except Exception as exc:  # noqa: BLE001
            cleanup_error = exc
            if published:
                try:
                    destination.unlink(missing_ok=True)
                except Exception:  # noqa: BLE001
                    pass
            try:
                shutil.rmtree(staging_dir)
            except Exception:  # noqa: BLE001
                pass
        if cleanup_error is not None:
            raise FormalBodySafetyError("DOCX 临时目录清理失败，已撤销发布") from cleanup_error


def _sanitize_markdown(markdown: str, removed: list[str]) -> str:
    original = str(markdown or "")
    removed_before = len(removed)
    output: list[str] = []
    for line_index, line in enumerate(original.splitlines()):
        output.append(_sanitize_text_segments(line, f"markdown.line[{line_index}]", removed))
    if len(removed) == removed_before:
        return original
    return re.sub(r"\n{3,}", "\n\n", "\n".join(output)).strip()


def _sanitize_report_ir(report_ir: dict[str, Any] | None, removed: list[str]) -> dict[str, Any] | None:
    if not isinstance(report_ir, dict):
        return None
    cleaned = copy.deepcopy(report_ir)
    for field in REPORT_IR_METADATA_FIELDS:
        if field in cleaned:
            cleaned[field] = _sanitize_scalar(cleaned.get(field), f"report_ir.{field}", removed)
    for field in ("lead_paragraphs", "enterprise_tips"):
        if isinstance(cleaned.get(field), list):
            cleaned[field] = _sanitize_string_list(cleaned[field], f"report_ir.{field}", removed)

    sections = cleaned.get("sections")
    if not isinstance(sections, list):
        return cleaned
    for section_index, section in enumerate(sections):
        if not isinstance(section, dict):
            continue
        section["heading"] = _sanitize_scalar(
            section.get("heading"), f"report_ir.sections[{section_index}].heading", removed
        )
        for field in ("paragraphs", "highlights"):
            if isinstance(section.get(field), list):
                section[field] = _sanitize_string_list(
                    section[field], f"report_ir.sections[{section_index}].{field}", removed
                )
        tables = section.get("tables")
        if not isinstance(tables, list):
            continue
        for table_index, table in enumerate(tables):
            if not isinstance(table, dict):
                continue
            prefix = f"report_ir.sections[{section_index}].tables[{table_index}]"
            table["title"] = _sanitize_scalar(table.get("title"), f"{prefix}.title", removed)
            for field in ("headers", "notes"):
                if isinstance(table.get(field), list):
                    table[field] = _sanitize_string_list(table[field], f"{prefix}.{field}", removed)
            rows = table.get("rows")
            if isinstance(rows, list):
                for row_index, row in enumerate(rows):
                    if not isinstance(row, list):
                        continue
                    rows[row_index] = [
                        _sanitize_scalar(value, f"{prefix}.rows[{row_index}][{column_index}]", removed)
                        for column_index, value in enumerate(row)
                    ]
    return cleaned


def _sanitize_scalar(value: Any, location: str, removed: list[str]) -> str:
    text = str(value or "")
    return _sanitize_text_segments(text, location, removed)


def _sanitize_text_segments(value: str, location: str, removed: list[str]) -> str:
    output: list[str] = []
    changed = False
    for segment_index, segment in enumerate(_SENTENCE_SPLIT_RE.split(value)):
        hits = find_forbidden_phrases(segment, f"{location}.sentence[{segment_index}]")
        if hits:
            changed = True
            removed.extend(hit.phrase for hit in hits)
            continue
        output.append(segment)
    return "".join(output).strip() if changed else value


def _sanitize_string_list(values: list[Any], prefix: str, removed: list[str]) -> list[str]:
    output: list[str] = []
    for index, value in enumerate(values):
        cleaned = _sanitize_scalar(value, f"{prefix}[{index}]", removed)
        if cleaned.strip():
            output.append(cleaned)
    return output


def _scan_paragraphs(paragraphs: Iterable[Any], prefix: str) -> list[ForbiddenPhraseHit]:
    hits: list[ForbiddenPhraseHit] = []
    for paragraph_index, paragraph in enumerate(paragraphs):
        hits.extend(find_forbidden_phrases(paragraph.text, f"{prefix}.paragraph[{paragraph_index}]"))
    return hits


def _scan_table(table: Any, prefix: str, depth: int) -> list[ForbiddenPhraseHit]:
    hits: list[ForbiddenPhraseHit] = []
    for row_index, row in enumerate(table.rows):
        for cell_index, cell in enumerate(row.cells):
            cell_prefix = f"{prefix}.row[{row_index}].cell[{cell_index}]"
            hits.extend(_scan_paragraphs(cell.paragraphs, cell_prefix))
            for nested_index, nested_table in enumerate(cell.tables):
                nested_prefix = f"{cell_prefix}.nested_table[{nested_index}]"
                hits.extend(_scan_table(nested_table, nested_prefix, depth + 1))
    return hits


def _dedupe_hits(hits: Iterable[ForbiddenPhraseHit]) -> tuple[ForbiddenPhraseHit, ...]:
    output: list[ForbiddenPhraseHit] = []
    seen: set[tuple[str, str]] = set()
    for hit in hits:
        key = (hit.phrase, hit.location)
        if key in seen:
            continue
        seen.add(key)
        output.append(hit)
    return tuple(output)
